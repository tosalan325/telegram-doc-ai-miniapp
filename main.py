import asyncio
import base64
import os
import uuid
from pathlib import Path
from typing import Optional

import pymupdf
import requests
from PIL import Image
from docx import Document
from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, Request, UploadFile
from fastapi.responses import JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pypdf import PdfReader

# =========================
# Настройки проекта
# =========================

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

OPENROUTER_API_KEY = os.getenv("OPENROUTER_API_KEY")
OPENROUTER_MODEL = os.getenv(
    "OPENROUTER_MODEL",
    "google/gemini-2.5-flash",
)
WEBAPP_URL = os.getenv(
    "WEBAPP_URL",
    "http://127.0.0.1:8000",
)

if not OPENROUTER_API_KEY:
    raise ValueError(
        "OPENROUTER_API_KEY не найден. Проверь файл .env"
    )

# =========================
# Папки
# =========================

DOWNLOADS_DIR = BASE_DIR / "downloads"
DOWNLOADS_DIR.mkdir(exist_ok=True)

# =========================
# FastAPI
# =========================

app = FastAPI(title="DocCheck Mini App")

app.mount(
    "/static",
    StaticFiles(directory=BASE_DIR / "static"),
    name="static",
)

templates = Jinja2Templates(
    directory=BASE_DIR / "templates"
)


@app.get("/")
async def index(request: Request):
    return templates.TemplateResponse(
        request=request,
        name="index.html",
        context={},
    )


# =========================
# Ограничения и форматы
# =========================

MAX_FILE_SIZE_MB = 15
MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024

ALLOWED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".jpg",
    ".jpeg",
    ".png",
    ".webp",
}

IMAGE_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".webp",
}


# =========================
# Извлечение текста
# =========================

def extract_text_from_pdf(file_path: Path) -> str:
    text_parts = []

    reader = PdfReader(str(file_path))

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text_parts.append(page_text)

    return "\n".join(text_parts).strip()


def extract_text_from_docx(file_path: Path) -> str:
    document = Document(str(file_path))
    text_parts = []

    # Обычные абзацы
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            text_parts.append(text)

    # Текст внутри таблиц
    for table in document.tables:
        for row in table.rows:
            row_parts = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    row_parts.append(cell_text)

            if row_parts:
                text_parts.append(" | ".join(row_parts))

    return "\n".join(text_parts).strip()


def extract_text_from_txt(file_path: Path) -> str:
    return file_path.read_text(
        encoding="utf-8",
        errors="ignore",
    ).strip()


def limit_text(
    text: str,
    max_chars: int = 25000,
) -> str:
    if len(text) <= max_chars:
        return text

    return (
        text[:max_chars]
        + "\n\n[Документ был сокращён из-за ограничения размера.]"
    )


# =========================
# Работа с изображениями
# =========================

def prepare_image_for_ai(input_path: Path) -> Path:
    output_path = (
        DOWNLOADS_DIR
        / f"{input_path.stem}_{uuid.uuid4().hex}.prepared.jpg"
    )

    with Image.open(input_path) as source_image:
        image = source_image.convert("RGB")

        max_side = 1600
        width, height = image.size

        if max(width, height) > max_side:
            if width >= height:
                new_width = max_side
                new_height = int(height * max_side / width)
            else:
                new_height = max_side
                new_width = int(width * max_side / height)

            image = image.resize(
                (new_width, new_height),
                Image.Resampling.LANCZOS,
            )

        image.save(
            output_path,
            format="JPEG",
            quality=85,
            optimize=True,
        )

    return output_path


def image_to_data_url(
    image_path: Path,
    generated_files: list[Path],
) -> str:
    prepared_path = prepare_image_for_ai(image_path)
    generated_files.append(prepared_path)

    with open(prepared_path, "rb") as image_file:
        encoded = base64.b64encode(
            image_file.read()
        ).decode("utf-8")

    return f"data:image/jpeg;base64,{encoded}"


def convert_pdf_pages_to_images(
    pdf_path: Path,
    generated_files: list[Path],
    max_pages: int = 5,
) -> list[Path]:
    image_paths = []

    pdf_document = pymupdf.open(str(pdf_path))

    try:
        pages_count = min(
            len(pdf_document),
            max_pages,
        )

        for page_index in range(pages_count):
            page = pdf_document[page_index]
            pix = page.get_pixmap(dpi=170)

            image_path = (
                DOWNLOADS_DIR
                / (
                    f"{pdf_path.stem}_"
                    f"{uuid.uuid4().hex}_"
                    f"page_{page_index + 1}.jpg"
                )
            )

            pix.save(str(image_path))

            image_paths.append(image_path)
            generated_files.append(image_path)
    finally:
        pdf_document.close()

    return image_paths


# =========================
# Системные промпты
# =========================

SYSTEM_PROMPT_TEXT = """
Ты — AI-помощник для краткого и точного анализа документов.

Обязательные правила:

1. Отвечай только по тем типам анализа, которые выбрал пользователь.
2. Не делай полный анализ, если пользователь не выбрал полный анализ.
3. Не добавляй разделы, которых пользователь не просил.
4. Не пересказывай весь документ без необходимости.
5. Не повторяй одну информацию в разных разделах.
6. Для каждого выбранного раздела используй не более 3–5 коротких пунктов.
7. Каждый пункт должен занимать не более 1–2 коротких предложений.
8. Если нужной информации нет, сообщи об этом одной фразой.
9. Не добавляй длинное вступление и длинное заключение.
10. Не выдумывай условия, которых нет в документе.
11. Пиши по-русски, простым и понятным языком.
12. Не принимай решение о подписании документа за пользователя.
13. Не утверждай, что документ законный или незаконный.
14. Если пользователь выбрал несколько тем, создай только соответствующие короткие разделы.
15. Не добавляй дисклеймер после каждого раздела.

В самом конце ответа добавь одну короткую строку:
«Это информационный разбор, а не юридическая консультация».
""".strip()


SYSTEM_PROMPT_IMAGE = """
Ты — AI-помощник для краткого и точного анализа документов по изображениям.

Обязательные правила:

1. Внимательно прочитай только видимый текст.
2. Отвечай только по выбранным пользователем темам.
3. Не делай полный анализ без прямого запроса.
4. Не добавляй лишние разделы.
5. Для каждого раздела используй не более 3–5 коротких пунктов.
6. Не повторяй содержание документа.
7. Не выдумывай текст, которого не видно.
8. Если часть изображения не читается, сообщи об этом кратко.
9. Пиши по-русски, простым и понятным языком.
10. Не принимай решение о подписании за пользователя.
11. Если страниц несколько, воспринимай их как один документ.

В самом конце ответа добавь одну короткую строку:
«Это информационный разбор, а не юридическая консультация».
""".strip()


# =========================
# Типы анализа
# =========================

ANALYSIS_TITLES = {
    "full": "🧾 Полный анализ",
    "summary": "📄 Краткое резюме",
    "strengths": "✅ Сильные стороны",
    "weaknesses": "⚠️ Слабые стороны",
    "risks": "❗ Риски документа",
    "dates": "📅 Даты и сроки",
    "money": "💰 Деньги и суммы",
    "parties": "👥 Стороны документа",
    "attention": "🔍 Важные пункты",
    "simple": "🧠 Простыми словами",
    "fixes": "✍️ Что исправить",
    "questions": "❓ Вопросы второй стороне",
    "dangerous": "🟥 Опасные формулировки",
    "score": "📊 Оценка документа",
    "custom": "💬 Ответ на вопрос",
}


ANALYSIS_PROMPTS = {
    "full": """
Сделай компактный общий анализ.

Формат:
1. Суть документа — не более 3 предложений.
2. Главные условия — не более 5 пунктов.
3. Возможные риски — не более 4 пунктов.
4. Что уточнить — не более 4 вопросов.

Не добавляй другие разделы.
""",

    "summary": """
Дай только краткое резюме документа.
Объясни его суть в 3–5 коротких предложениях.
Не добавляй отдельный анализ рисков, денег и сроков.
""",

    "strengths": """
Покажи только сильные стороны документа.
Укажи не более 5 коротких пунктов.
Не добавляй слабые стороны, риски и общий пересказ.
""",

    "weaknesses": """
Покажи только слабые стороны документа.
Укажи не более 5 коротких пунктов.
Обращай внимание на неясные, неполные или отсутствующие условия.
""",

    "risks": """
Покажи только возможные риски.
Укажи не более 5 пунктов.

Для каждого риска напиши:
- короткое название;
- уровень: низкий, средний или высокий;
- одно короткое объяснение.

Не пересказывай документ полностью.
""",

    "dates": """
Найди только даты и сроки.

Для каждого пункта укажи:
- дату или срок;
- что он означает.

Если дат или сроков нет, сообщи об этом одной фразой.
""",

    "money": """
Найди только финансовые условия:
- суммы;
- порядок оплаты;
- авансы и залоги;
- штрафы, пени и комиссии;
- условия возврата.

Не более 6 коротких пунктов.
""",

    "parties": """
Укажи только:
- стороны документа;
- роль каждой стороны;
- главные обязанности каждой стороны.

Не более 6 коротких пунктов.
""",

    "attention": """
Покажи только самые важные условия, которые легко пропустить.
Не более 5 коротких пунктов.
Для каждого пункта одной фразой объясни, почему он важен.
""",

    "simple": """
Объясни документ максимально простыми словами.
Не более 5 коротких предложений.
Не используй сложные юридические формулировки.
""",

    "fixes": """
Покажи только то, что желательно исправить, добавить или уточнить.
Не более 5 конкретных рекомендаций.
Не переписывай весь документ.
""",

    "questions": """
Составь только список вопросов второй стороне.
Не более 7 коротких и конкретных вопросов.
Не добавляй общий анализ документа.
""",

    "dangerous": """
Найди только потенциально опасные, спорные или размытые формулировки.
Не более 5 пунктов.

Для каждого пункта укажи:
- короткую цитату или точный пересказ;
- почему это стоит уточнить.
""",

    "score": """
Дай только краткую ориентировочную оценку:

- понятность: от 1 до 10;
- полнота условий: от 1 до 10;
- финансовая прозрачность: от 1 до 10;
- общий риск: низкий, средний или высокий;
- одна главная рекомендация.

Не добавляй полный анализ.
""",

    "custom": """
Ответь только на вопрос пользователя.
Ответ должен занимать не более 5 коротких предложений.

Используй только информацию из документа.
Если ответа в документе нет, прямо сообщи об этом.
""",
}


# =========================
# Обработка выбранных типов
# =========================

def parse_analysis_types(
    raw_value: str,
) -> list[str]:
    requested_types = [
        item.strip()
        for item in raw_value.split(",")
        if item.strip()
    ]

    # Сравнение всегда является отдельным режимом
    if "compare" in requested_types:
        return ["compare"]

    valid_types = []
    seen = set()

    for analysis_type in requested_types:
        if (
            analysis_type in ANALYSIS_PROMPTS
            and analysis_type not in seen
        ):
            valid_types.append(analysis_type)
            seen.add(analysis_type)

    if not valid_types:
        return ["full"]

    # Полный анализ нельзя объединять с другими режимами
    if "full" in valid_types:
        return ["full"]

    return valid_types


def get_role_instruction(
    user_role: str,
) -> str:
    if user_role == "author":
        return """
Пользователь сам составляет или редактирует документ.

Учитывай эту ситуацию только внутри выбранных задач:
- обращай внимание на недостающие условия;
- отмечай размытые формулировки;
- предлагай, что стоит уточнить;
- не добавляй отдельный общий анализ без запроса.
""".strip()

    return """
Пользователь проверяет документ перед подписанием.

Учитывай эту ситуацию только внутри выбранных задач:
- оценивай условия с позиции подписанта;
- обращай внимание на деньги, сроки и ответственность;
- не добавляй отдельный общий анализ без запроса.
""".strip()


def build_analysis_instruction(
    analysis_types: list[str],
    custom_question: str = "",
) -> str:
    parts = []

    for analysis_type in analysis_types:
        instruction = ANALYSIS_PROMPTS.get(
            analysis_type
        )

        if not instruction:
            continue

        title = ANALYSIS_TITLES.get(
            analysis_type,
            analysis_type,
        )

        custom_part = ""

        if analysis_type == "custom":
            custom_part = (
                "\n\nВопрос пользователя:\n"
                f"{custom_question.strip()}"
            )

        parts.append(
            f"""
РАЗДЕЛ: {title}

{instruction.strip()}
{custom_part}
""".strip()
        )

    joined_parts = "\n\n".join(parts)

    return f"""
Выполни только перечисленные ниже задачи.

Не добавляй никаких других разделов.
Не превращай ответ в полный анализ.
Не повторяй одну и ту же информацию.
Не пиши длинное вступление.
Сохраняй указанные ограничения по количеству пунктов.

{joined_parts}
""".strip()


def calculate_max_tokens(
    analysis_types: list[str],
) -> int:
    if analysis_types == ["full"]:
        return 1300

    number_of_types = len(analysis_types)

    if number_of_types == 1:
        return 600

    return min(
        1800,
        350 + number_of_types * 240,
    )


# =========================
# OpenRouter
# =========================

def call_openrouter(
    payload: dict,
    timeout: int = 240,
) -> str:
    url = (
        "https://openrouter.ai/api/v1/"
        "chat/completions"
    )

    headers = {
        "Authorization": (
            f"Bearer {OPENROUTER_API_KEY}"
        ),
        "Content-Type": "application/json",
        "HTTP-Referer": WEBAPP_URL,
        "X-Title": "DocCheck Mini App",
    }

    response = requests.post(
        url,
        headers=headers,
        json=payload,
        timeout=timeout,
    )

    if response.status_code != 200:
        raise Exception(
            "OpenRouter API error "
            f"{response.status_code}: "
            f"{response.text}"
        )

    data = response.json()

    try:
        content = data["choices"][0]["message"]["content"]

        if isinstance(content, str):
            return content.strip()

        raise ValueError(
            "Содержимое ответа не является строкой."
        )
    except Exception as error:
        raise Exception(
            f"Неожиданный ответ от OpenRouter: {data}"
        ) from error


# =========================
# Анализ текста
# =========================

def analyze_document_text_with_ai(
    document_text: str,
    analysis_types: list[str],
    user_role: str = "signer",
    custom_question: str = "",
) -> str:
    document_text = limit_text(document_text)

    role_instruction = get_role_instruction(
        user_role
    )

    analysis_instruction = (
        build_analysis_instruction(
            analysis_types,
            custom_question,
        )
    )

    user_prompt = f"""
Контекст пользователя:

{role_instruction}

Выбранные задачи:

{analysis_instruction}

Текст документа:

{document_text}
""".strip()

    payload = {
        "model": OPENROUTER_MODEL,
        "messages": [
            {
                "role": "system",
                "content": SYSTEM_PROMPT_TEXT,
            },
            {
                "role": "user",
                "content": user_prompt,
            },
        ],
        "temperature": 0.1,
        "max_tokens": calculate_max_tokens(
            analysis_types
        ),
    }

    return call_openrouter(payload)


# =========================
# Анализ изображений
# =========================

def analyze_document_images_with_ai(
    image_paths: list[Path],
    analysis_types: list[str],
    user_role: str = "signer",
    custom_question: str = "",
) -> str:
    generated_files: list[Path] = []

    try:
        role_instruction = get_role_instruction(
            user_role
        )

        analysis_instruction = (
            build_analysis_instruction(
                analysis_types,
                custom_question,
            )
        )

        content = [
            {
                "type": "text",
                "text": f"""
На изображении или изображениях находится один документ.

Прочитай видимый текст.
Если страниц несколько, воспринимай их как один документ.
Если часть текста не читается, кратко сообщи об этом.
Не выдумывай отсутствующую информацию.

Контекст пользователя:

{role_instruction}

Выбранные задачи:

{analysis_instruction}
""".strip(),
            }
        ]

        for image_path in image_paths:
            image_data_url = image_to_data_url(
                image_path,
                generated_files,
            )

            content.append(
                {
                    "type": "image_url",
                    "image_url": {
                        "url": image_data_url,
                    },
                }
            )

        payload = {
            "model": OPENROUTER_MODEL,
            "messages": [
                {
                    "role": "system",
                    "content": SYSTEM_PROMPT_IMAGE,
                },
                {
                    "role": "user",
                    "content": content,
                },
            ],
            "temperature": 0.1,
            "max_tokens": calculate_max_tokens(
                analysis_types
            ),
        }

        return call_openrouter(
            payload,
            timeout=240,
        )
    finally:
        delete_files(generated_files)


# =========================
# Сохранение файлов
# =========================

async def save_upload_file(
    file: UploadFile,
) -> tuple[Path, str]:
    original_name = Path(
        file.filename or "document"
    ).name

    file_extension = Path(
        original_name
    ).suffix.lower()

    if file_extension not in ALLOWED_EXTENSIONS:
        raise ValueError(
            "Неподдерживаемый формат. "
            "Можно загрузить PDF, DOCX, TXT, "
            "JPG, PNG или WEBP."
        )

    file_bytes = await file.read()

    if not file_bytes:
        raise ValueError("Загруженный файл пустой.")

    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        raise ValueError(
            "Файл слишком большой. "
            f"Максимальный размер: {MAX_FILE_SIZE_MB} МБ."
        )

    safe_name = (
        original_name
        .replace(" ", "_")
        .replace("/", "_")
        .replace("\\", "_")
    )

    local_file_path = (
        DOWNLOADS_DIR
        / f"{uuid.uuid4().hex}_{safe_name}"
    )

    with open(local_file_path, "wb") as output_file:
        output_file.write(file_bytes)

    return local_file_path, file_extension


# =========================
# Подготовка содержимого
# =========================

def extract_document_content(
    local_file_path: Path,
    file_extension: str,
    generated_files: list[Path],
) -> dict:
    if file_extension in IMAGE_EXTENSIONS:
        return {
            "kind": "images",
            "images": [local_file_path],
        }

    if file_extension == ".txt":
        extracted_text = extract_text_from_txt(
            local_file_path
        )

        if not extracted_text:
            raise ValueError(
                "TXT-файл пустой или текст "
                "не удалось прочитать."
            )

        return {
            "kind": "text",
            "text": extracted_text,
        }

    if file_extension == ".docx":
        extracted_text = extract_text_from_docx(
            local_file_path
        )

        if not extracted_text:
            raise ValueError(
                "В DOCX не найден редактируемый текст. "
                "Возможно, документ содержит только скан."
            )

        return {
            "kind": "text",
            "text": extracted_text,
        }

    if file_extension == ".pdf":
        extracted_text = extract_text_from_pdf(
            local_file_path
        )

        if extracted_text:
            return {
                "kind": "text",
                "text": extracted_text,
            }

        image_paths = (
            convert_pdf_pages_to_images(
                local_file_path,
                generated_files,
                max_pages=5,
            )
        )

        if not image_paths:
            raise ValueError(
                "PDF не содержит текста, "
                "и его не удалось преобразовать "
                "в изображения."
            )

        return {
            "kind": "images",
            "images": image_paths,
        }

    raise ValueError(
        "Файл не удалось обработать."
    )


# =========================
# Сравнение документов
# =========================

def compare_documents_with_ai(
    first_content: dict,
    second_content: dict,
    user_role: str = "signer",
) -> str:
    role_instruction = get_role_instruction(
        user_role
    )

    compare_instruction = f"""
{role_instruction}

Кратко сравни два документа.

Покажи только:

1. Главные изменения — не более 5 пунктов.
2. Изменения в деньгах — не более 3 пунктов.
3. Изменения в сроках — не более 3 пунктов.
4. Что стало рискованнее — не более 3 пунктов.
5. Что стало лучше — не более 3 пунктов.

Не пересказывай документы полностью.
Не повторяй одинаковую информацию.
Если значимых различий нет, сообщи об этом кратко.
""".strip()

    # Оба документа содержат обычный текст
    if (
        first_content["kind"] == "text"
        and second_content["kind"] == "text"
    ):
        user_prompt = f"""
{compare_instruction}

ДОКУМЕНТ 1:

{limit_text(first_content["text"], 16000)}

ДОКУМЕНТ 2:

{limit_text(second_content["text"], 16000)}
""".strip()

        payload = {
            "model": OPENROUTER_MODEL,
            "messages": [
                {
                    "role": "system",
                    "content": SYSTEM_PROMPT_TEXT,
                },
                {
                    "role": "user",
                    "content": user_prompt,
                },
            ],
            "temperature": 0.1,
            "max_tokens": 1200,
        }

        return call_openrouter(
            payload,
            timeout=240,
        )

    generated_files: list[Path] = []

    try:
        content = [
            {
                "type": "text",
                "text": compare_instruction,
            }
        ]

        append_document_to_image_content(
            content=content,
            document_content=first_content,
            document_number=1,
            generated_files=generated_files,
        )

        append_document_to_image_content(
            content=content,
            document_content=second_content,
            document_number=2,
            generated_files=generated_files,
        )

        payload = {
            "model": OPENROUTER_MODEL,
            "messages": [
                {
                    "role": "system",
                    "content": SYSTEM_PROMPT_IMAGE,
                },
                {
                    "role": "user",
                    "content": content,
                },
            ],
            "temperature": 0.1,
            "max_tokens": 1200,
        }

        return call_openrouter(
            payload,
            timeout=240,
        )
    finally:
        delete_files(generated_files)


def append_document_to_image_content(
    content: list,
    document_content: dict,
    document_number: int,
    generated_files: list[Path],
) -> None:
    if document_content["kind"] == "text":
        content.append(
            {
                "type": "text",
                "text": (
                    f"ДОКУМЕНТ {document_number}, текст:\n\n"
                    f"{limit_text(document_content['text'], 14000)}"
                ),
            }
        )

        return

    content.append(
        {
            "type": "text",
            "text": (
                f"ДОКУМЕНТ {document_number}, "
                "изображения:"
            ),
        }
    )

    for image_path in document_content["images"]:
        content.append(
            {
                "type": "image_url",
                "image_url": {
                    "url": image_to_data_url(
                        image_path,
                        generated_files,
                    ),
                },
            }
        )


# =========================
# Удаление временных файлов
# =========================

def delete_files(
    file_paths: list[Path],
) -> None:
    unique_paths = set(file_paths)

    for file_path in unique_paths:
        try:
            if file_path and file_path.exists():
                file_path.unlink()
        except OSError:
            pass


# =========================
# API анализа
# =========================

@app.post("/api/analyze")
async def analyze_file(
    file: UploadFile = File(...),
    file2: Optional[UploadFile] = File(None),
    analysis_types: str = Form("full"),
    user_role: str = Form("signer"),
    custom_question: str = Form(""),
):
    saved_files: list[Path] = []
    generated_files: list[Path] = []

    try:
        if user_role not in {"signer", "author"}:
            user_role = "signer"

        selected_analysis_types = (
            parse_analysis_types(
                analysis_types
            )
        )

        if (
            "custom" in selected_analysis_types
            and not custom_question.strip()
        ):
            return JSONResponse(
                status_code=400,
                content={
                    "ok": False,
                    "error": (
                        "Введите свой вопрос "
                        "по документу."
                    ),
                },
            )

        first_file_path, first_extension = (
            await save_upload_file(file)
        )

        saved_files.append(first_file_path)

        first_content = extract_document_content(
            first_file_path,
            first_extension,
            generated_files,
        )

        # Сравнение двух документов
        if selected_analysis_types == ["compare"]:
            if file2 is None or not file2.filename:
                return JSONResponse(
                    status_code=400,
                    content={
                        "ok": False,
                        "error": (
                            "Для сравнения нужно "
                            "загрузить второй документ."
                        ),
                    },
                )

            second_file_path, second_extension = (
                await save_upload_file(file2)
            )

            saved_files.append(second_file_path)

            second_content = (
                extract_document_content(
                    second_file_path,
                    second_extension,
                    generated_files,
                )
            )

            result = await asyncio.to_thread(
                compare_documents_with_ai,
                first_content,
                second_content,
                user_role,
            )

            return {
                "ok": True,
                "analysis_types": ["compare"],
                "result": result,
            }

        # Анализ текстового документа
        if first_content["kind"] == "text":
            result = await asyncio.to_thread(
                analyze_document_text_with_ai,
                first_content["text"],
                selected_analysis_types,
                user_role,
                custom_question.strip(),
            )

            return {
                "ok": True,
                "analysis_types": (
                    selected_analysis_types
                ),
                "result": result,
            }

        # Анализ изображений
        if first_content["kind"] == "images":
            result = await asyncio.to_thread(
                analyze_document_images_with_ai,
                first_content["images"],
                selected_analysis_types,
                user_role,
                custom_question.strip(),
            )

            return {
                "ok": True,
                "analysis_types": (
                    selected_analysis_types
                ),
                "result": result,
            }

        return JSONResponse(
            status_code=400,
            content={
                "ok": False,
                "error": (
                    "Файл не удалось обработать."
                ),
            },
        )

    except ValueError as error:
        return JSONResponse(
            status_code=400,
            content={
                "ok": False,
                "error": str(error),
            },
        )

    except requests.Timeout:
        return JSONResponse(
            status_code=504,
            content={
                "ok": False,
                "error": (
                    "AI не успел обработать документ. "
                    "Попробуйте ещё раз."
                ),
            },
        )

    except Exception as error:
        return JSONResponse(
            status_code=500,
            content={
                "ok": False,
                "error": str(error),
            },
        )

    finally:
        delete_files(generated_files)
        delete_files(saved_files)

        try:
            await file.close()
        except Exception:
            pass

        if file2 is not None:
            try:
                await file2.close()
            except Exception:
                pass